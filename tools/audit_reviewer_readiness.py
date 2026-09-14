"""Independent checks of the distributed scientific archive and manuscript.

No primary metric or bootstrap function is imported. Optional model refits use
the released candidate builders, with fresh estimators and no checkpoints.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import time
from zipfile import ZipFile

import numpy as np
import pandas as pd
from sklearn.model_selection import GroupKFold, KFold

ROOT = Path(__file__).resolve().parents[1]


def check(condition, label):
    if not condition:
        raise RuntimeError(label)


def close(actual, expected, label, tolerance=1e-10):
    check(np.allclose(actual, expected, atol=tolerance, rtol=tolerance), label)


def metrics(truth, prediction):
    truth, prediction = np.asarray(truth, float), np.asarray(prediction, float)
    error = prediction - truth
    return dict(mae=float(np.mean(np.abs(error))), rmse=float(np.sqrt(np.mean(error**2))),
                r2=float(1 - np.sum(error**2) / np.sum((truth-truth.mean())**2)) if np.ptp(truth) else None,
                within_1_day=float(np.mean(np.abs(error) <= 1)), bias=float(error.mean()))


def main(source_zip=None, refit_fold=None, output=None):
    started = time.monotonic()
    data_path = ROOT/'01_DATASET/barley_nir_long.npz'
    digest = hashlib.sha256(data_path.read_bytes()).hexdigest()
    check(digest == '3ee66476e6564e6b20d6b3c87433517e04fb6ce99e8ea1625e3f8cc1a3c70cdd', 'Prepared dataset digest')
    z = np.load(data_path, allow_pickle=True)
    y, dish, grain = z['day'], z['petri_dish'], z['grain_id']
    uid = np.array([f'{d}_{g}' for d,g in zip(dish,grain)])
    results = ROOT/'03_RESULTADOS'
    run = json.loads((results/'nested_run_metadata.json').read_text())
    check(hashlib.sha256((ROOT/'02_CODIGO_NOTEBOOKS/run_nested_validation.py').read_bytes()).hexdigest()
          == run['released_script_sha256'], 'Released primary script digest')
    oof = pd.read_csv(results/'nested_oof_predictions.csv.gz')
    outer = pd.read_csv(results/'nested_outer_fold_metrics.csv')
    inner = pd.read_csv(results/'nested_inner_search.csv.gz')
    outer_splits = list(GroupKFold(5).split(y,y,dish))
    checks = []
    for fold,(train,test) in enumerate(outer_splits,1):
        for family, count in [('PLSR',9),('SVR_RBF',8),('ExtraTrees',6)]:
            selection = inner.loc[inner.outer_fold.eq(fold)&inner.family.eq(family)]
            check(len(selection)==count*4 and not selection.duplicated(['candidate_key','inner_fold']).any(), 'Inner candidate coverage')
            local_splits = list(GroupKFold(4).split(y[train],y[train],dish[train]))
            for k,(tr,va) in enumerate(local_splits,1):
                block=selection.loc[selection.inner_fold.eq(k)]
                check(block.n_train.eq(len(tr)).all() and block.n_valid.eq(len(va)).all(), 'Inner partition sizes')
                check(set(dish[train[tr]]).isdisjoint(dish[train[va]]), 'Inner dish overlap')
            ranking=selection.groupby('candidate_key').agg(mae=('mae','mean'),rmse=('rmse','mean')).reset_index().sort_values(['mae','rmse','candidate_key'])
            saved=outer.loc[outer.outer_fold.eq(fold)&outer.family.eq(family)].iloc[0]
            check(saved.candidate_key==ranking.iloc[0].candidate_key, 'Independent candidate ranking')
            check(selection.selected.eq(selection.candidate_key.eq(saved.candidate_key)).all(), 'Selected flags')
            values=oof.loc[oof.outer_fold.eq(fold)&oof.family.eq(family)].sort_values('row_index')
            check(np.array_equal(values.row_index,test), 'Primary row membership')
            check(np.array_equal(values.grain_uid,uid[test]), 'OOF kernel metadata')
            check(np.array_equal(values.cultivar,z['cultivar'][test]) and np.array_equal(values.sample_group,z['sample_group'][test]), 'OOF group metadata')
            for name,value in metrics(y[test],values.y_pred).items():close(value,saved[name],'Outer metric '+name)
    checks.append('Complete inner memberships, rankings, selection flags and primary metadata')

    # Independently reconstruct the exact paired draw sequence using cluster sums.
    boot=pd.read_csv(results/'nested_cluster_bootstrap_replicates.csv.gz').set_index(['family','bootstrap_repetition'])
    dishes=np.array(sorted(np.unique(dish)))
    rng=np.random.default_rng(2026)
    weights=np.array([np.bincount(np.searchsorted(dishes,rng.choice(dishes,len(dishes),replace=True)),minlength=len(dishes)) for _ in range(2000)])
    maximum=0.0
    for family,frame in oof.groupby('family'):
        blocks=[]
        for d in dishes:
            b=frame.loc[frame.petri_dish.eq(d)];t=b.y_true.to_numpy(float);e=b.y_pred.to_numpy(float)-t
            blocks.append([len(b),np.abs(e).sum(),(e**2).sum(),e.sum(),(np.abs(e)<=1).sum(),t.sum(),(t**2).sum()])
        sums=weights@np.array(blocks);n,absolute,squared,bias,within,total,total2=sums.T
        recalculated=dict(mae=absolute/n,rmse=np.sqrt(squared/n),r2=1-squared/(total2-total**2/n),bias=bias/n,within_1_day=within/n)
        saved=boot.loc[family].sort_index()
        for name,value in recalculated.items():
            maximum=max(maximum,float(np.max(np.abs(value-saved[name].to_numpy()))))
            close(value,saved[name],'Regenerated paired bootstrap '+name)
    checks.append('All 6000 paired bootstrap replicates regenerated from seed and cluster sums')

    schemes={'random_row':list(KFold(5,shuffle=True,random_state=2026).split(y)),
             'grain_disjoint':list(GroupKFold(5).split(y,y,uid)),
             'dish_disjoint':outer_splits}
    fixed=pd.read_csv(results/'split_sensitivity_oof_predictions.csv.gz')
    fixed_metrics=pd.read_csv(results/'split_sensitivity_fold_metrics.csv')
    fixed_summary=pd.read_csv(results/'split_sensitivity_summary.csv')
    overlaps={}
    for scheme,splits in schemes.items():
        records=[];grain_overlap=[];dish_overlap=[]
        for fold,(train,test) in enumerate(splits,1):
            b=fixed.loc[fixed.scheme.eq(scheme)&fixed.fold.eq(fold)].sort_values('row_index')
            check(np.array_equal(b.row_index,test) and np.array_equal(b.y_true,y[test]),'Sensitivity membership')
            score=metrics(y[test],b.y_pred);records.append(score)
            expected=fixed_metrics.loc[fixed_metrics.scheme.eq(scheme)&fixed_metrics.fold.eq(fold)].iloc[0]
            for name,value in score.items():close(value,expected[name],'Sensitivity fold metric')
            grain_overlap.append(len(set(uid[train])&set(uid[test])))
            dish_overlap.append(len(set(dish[train])&set(dish[test])))
        expected=fixed_summary.loc[fixed_summary.scheme.eq(scheme)].iloc[0]
        for name in records[0]:close(np.mean([r[name] for r in records]),expected[name],'Sensitivity mean metric')
        overlaps[scheme]={'kernel_overlap_each_fold':grain_overlap,'dish_overlap_each_fold':dish_overlap}
    check(not any(overlaps['grain_disjoint']['kernel_overlap_each_fold']), 'Kernel split overlap')
    check(not any(overlaps['dish_disjoint']['dish_overlap_each_fold']), 'Dish split overlap')
    checks.append('Fixed-partition memberships, overlaps and all reported means')

    domain_inner=pd.read_csv(results/'domain_stress_inner_search.csv.gz')
    domains=pd.read_csv(results/'domain_stress_test_summary.csv')
    for row in domains.itertuples():
        train=np.flatnonzero(z['cultivar' if row.scheme=='leave_one_cultivar_out' else 'sample_group']!=row.held_out)
        selection=domain_inner.loc[domain_inner.scheme.eq(row.scheme)&domain_inner.held_out.eq(row.held_out)]
        check(len(selection)==32 and not selection.duplicated(['candidate_key','inner_fold']).any(), 'Domain candidate coverage')
        for fold,(tr,va) in enumerate(GroupKFold(4).split(y[train],y[train],dish[train]),1):
            b=selection.loc[selection.inner_fold.eq(fold)]
            check(b.n_train.eq(len(tr)).all() and b.n_valid.eq(len(va)).all(), 'Domain inner memberships')
        rank=selection.groupby('candidate_key').agg(mae=('mae','mean'),rmse=('rmse','mean')).reset_index().sort_values(['mae','rmse','candidate_key'])
        check(row.preprocessing+'|'+row.parameter==rank.iloc[0].candidate_key,'Domain selection ranking')
    checks.append('All six domain inner selections independently checked')

    # Compare every numeric cell in the manuscript's five scientific tables.
    from docx import Document
    manuscript=ROOT/'04_MANUSCRITO_CCIS/ARTICULO_CEBADA_NIR_GITHUB.docx'
    doc=Document(manuscript);check(len(doc.tables)==5 and len(doc.inline_shapes)==3,'Five tables and three figures')
    check([len(t.rows) for t in doc.tables]==[4,7,4,7,4],'Complete manuscript table rows')
    ci=pd.read_csv(results/'nested_cluster_bootstrap_ci.csv')
    for cells,family in zip(list(doc.tables[0].rows)[1:],['SVR_RBF','PLSR','ExtraTrees']):
        b=oof.loc[oof.family.eq(family)];score=metrics(b.y_true,b.y_pred);interval=ci.loc[ci.family.eq(family)&ci.metric.eq('mae')].iloc[0]
        check(cells.cells[1].text==f"{score['mae']:.3f} [{interval.ci_low:.3f}, {interval.ci_high:.3f}]",'Manuscript MAE interval rounding')
        for c,v,decimals in zip(cells.cells[2:],[score['rmse'],score['r2'],100*score['within_1_day']],[3,3,1]):check(c.text==f'{v:.{decimals}f}','Manuscript table 1 rounding')
    perday=pd.read_csv(results/'best_model_metrics_by_day.csv')
    for cells,row in zip(list(doc.tables[1].rows)[1:],perday.itertuples()):
        expected=[str(int(row.day_since_moisture)),str(row.n),f'{row.mae:.3f}',f'{row.rmse:.3f}',f'{row.bias:.3f}',f'{100*row.within_1_day:.1f}']
        check([c.text for c in cells.cells]==expected,'Manuscript table 2')
        b=oof.loc[oof.family.eq('SVR_RBF')&oof.y_true.eq(row.day_since_moisture)]
        for name,value in metrics(b.y_true,b.y_pred).items():
            if name!='r2':close(value,getattr(row,name),'Stage metric')
    for cells,scheme in zip(list(doc.tables[2].rows)[1:],['random_row','grain_disjoint','dish_disjoint']):
        row=fixed_summary.loc[fixed_summary.scheme.eq(scheme)].iloc[0]
        check([c.text for c in cells.cells[1:]]==[f'{row.mae:.3f}',f'{row.rmse:.3f}',f'{row.r2:.3f}',f'{100*row.within_1_day:.1f}'],'Manuscript table 3')
    for cells,row in zip(list(doc.tables[3].rows)[1:],domains.itertuples()):
        check([c.text for c in cells.cells[1:]]==[row.held_out,f'{row.n_train_dishes}/{row.n_test_dishes}',f'{row.mae:.3f}',f'{row.rmse:.3f}',f'{row.r2:.3f}'],'Manuscript table 4')
    frozen=pd.read_csv(results/'revision_r1/svr_frozen_C_summary.csv')
    for cells,c in zip(list(doc.tables[4].rows)[1:],[100,300,1000]):
        a=frozen.loc[frozen.C.eq(c)&frozen.subset.eq('all_six_stages')].iloc[0];b=frozen.loc[frozen.C.eq(c)&frozen.subset.eq('post_moisture_1_to_5')].iloc[0]
        check([v.text for v in cells.cells]==[str(c),f'{a.mae:.3f}',f'{a.rmse:.3f}',f'{b.mae:.3f}',f'{b.r2:.3f}'],'Manuscript table 5')
    text='\n'.join(p.text for p in doc.paragraphs)
    check('https://github.com/Julians30/barley-nir-stage-benchmark' in text,'Manuscript repository link')
    check(not re.search(r'to be inserted|author\(s\)|authors to confirm|\[insert',text,re.I),'Manuscript placeholders')
    with ZipFile(manuscript) as archive:
        embedded={hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
    for file in ['fig_best_observed_vs_predicted.png','revision_r1/fig_methodology_workflow.png','revision_r1/fig_spectral_interpretation.png']:
        check(hashlib.sha256((results/file).read_bytes()).hexdigest() in embedded,'Embedded scientific figure '+file)
    checks.append('Every numeric manuscript table cell, three original embedded figures and repository link')

    source_report=None
    if source_zip:
        source_zip=Path(source_zip)
        with ZipFile(source_zip) as archive:
            names=archive.namelist();annotation_name=next(n for n in names if n.endswith('/annotations.csv'))
            annotations=pd.read_csv(io.BytesIO(archive.read(annotation_name))).set_index(['petri_dish','grain_id'])
            for stage in range(6):
                label='pre_moisture' if stage==0 else str(stage)
                matrix_name=next(n for n in names if n.endswith('/'+label+'/mean_spectra.npy'))
                order_name=next(n for n in names if n.endswith('/'+label+'/mean_spectra_order.csv'))
                matrix=np.load(io.BytesIO(archive.read(matrix_name)),allow_pickle=False)
                order=pd.read_csv(io.BytesIO(archive.read(order_name)));keys=pd.MultiIndex.from_frame(order[['petri_dish','grain_id']])
                prepared=np.flatnonzero(y==stage)
                positions={key:i for i,key in enumerate(keys)}
                aligned=np.array([positions[(d,g)] for d,g in zip(dish[prepared],grain[prepared])])
                check(np.array_equal(matrix[aligned],z['X'][prepared]),'All source spectrum values')
                annotation=annotations.loc[pd.MultiIndex.from_arrays([dish[prepared],grain[prepared]])]
                check(np.array_equal(annotation.germination_day,z['germination_day'][prepared]),'Source germination annotations')
                check(np.array_equal(annotation.variety,z['sample_group'][prepared]),'Source sample-group annotations')
        source_report={'archive_sha256':hashlib.sha256(source_zip.read_bytes()).hexdigest(),'spectra_exactly_matched':len(y),'channels_each':204,'germinating_kernels':int(np.sum(z['germination_day'][y==0]>=1))}
        checks.append('Every prepared spectral value and original annotation matched the retained source subset')

    refits=[]
    if refit_fold:
        sys.path.insert(0,str(ROOT/'02_CODIGO_NOTEBOOKS'))
        import run_nested_validation as primary
        from threadpoolctl import threadpool_limits
        variants=primary.spectral_variants(z['X'].astype(np.float32))
        train,test=outer_splits[refit_fold-1]
        for family,candidates in primary.candidate_grid().items():
            row=outer.loc[outer.outer_fold.eq(refit_fold)&outer.family.eq(family)].iloc[0]
            candidate=next(c for c in candidates if c.key==row.candidate_key);estimator=candidate.build()
            if family=='ExtraTrees':estimator.set_params(n_jobs=2)
            before=time.monotonic()
            with threadpool_limits(limits=2):
                estimator.fit(variants[candidate.preprocessing][train],y[train].astype(np.float32))
                prediction=np.asarray(estimator.predict(variants[candidate.preprocessing][test])).reshape(-1)
            saved=oof.loc[oof.outer_fold.eq(refit_fold)&oof.family.eq(family)].sort_values('row_index')
            difference=float(np.max(np.abs(prediction-saved.y_pred)))
            close(prediction,saved.y_pred,'Fresh selected outer fit '+family,1e-7)
            refits.append({'family':family,'outer_fold':refit_fold,'max_absolute_prediction_difference':difference,'fit_predict_seconds':time.monotonic()-before})
        checks.append('Fresh checkpoint-free outer fit for each family in one selected fold')
    hashes=[hashlib.sha256(np.ascontiguousarray(row).tobytes()).hexdigest() for row in z['X']]
    duplicate_count=len(hashes)-len(set(hashes))
    report={'result':'passed','audited_base_commit':'9c998bd2ec5e3bbe1119e8247bdcbf6e9cbf1d2e','checks':checks,
            'bootstrap_max_absolute_difference':maximum,'exact_duplicate_spectra':duplicate_count,
            'partition_overlap_counts':overlaps,'source_subset_comparison':source_report,'fresh_outer_refits':refits,
            'full_nested_search_repeated':False,'elapsed_seconds':time.monotonic()-started,
            'scope':'Independent saved-result calculations, seed regeneration, manuscript checks; optional retained-source comparison and selected-fold refits. No independent acquisition sequence or full nested retraining.'}
    if output:Path(output).write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-zip',type=Path,help='Optional retained compact original subset ZIP; not distributed in this repository.')
    parser.add_argument('--refit-fold',type=int,choices=range(1,6),help='Freshly refit all three selected estimators on this outer fold; does not repeat inner tuning.')
    parser.add_argument('--output',type=Path)
    args=parser.parse_args();main(args.source_zip,args.refit_fold,args.output)
